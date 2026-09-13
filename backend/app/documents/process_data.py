from __future__ import annotations

import copy
import unicodedata
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.documents.repository import DocumentRepository

OUTCOMES_SHEET = "Resultados dos processos"
SUBSIDIES_SHEET = "Subsídios disponibilizados"
PROCESS_COLUMN = "Número do processo"
SUBSIDY_PROCESS_COLUMN = "Número do processos"
BINARY_COLUMNS = {
    "contract": "Contrato",
    "bank_statement": "Extrato",
    "credit_proof": "Comprovante de crédito",
    "dossier": "Dossiê",
    "debt_evolution": "Demonstrativo de evolução da dívida",
    "referenced_report": "Laudo referenciado",
}
EXCLUDED_POST_OUTCOME_COLUMNS = (
    "Resultado macro",
    "Resultado micro",
    "Valor da condenação/indenização",
)


def canonical_process_number(value: str) -> str:
    canonical = "".join(character for character in value if character.isdigit())
    if not canonical:
        raise ValueError("Process number must contain digits")
    return canonical


class ProcessDataRepository:
    def __init__(self, documents: DocumentRepository) -> None:
        self.documents = documents

    def find(self, workbook_path: str, process_number: str) -> dict[str, Any]:
        path, normalized = self.documents.resolve(
            workbook_path,
            expected_suffixes={".xlsx", ".xlsm"},
        )
        stat = path.stat()
        record = self._find_cached(
            str(path),
            stat.st_mtime_ns,
            stat.st_size,
            canonical_process_number(process_number),
        )
        result = copy.deepcopy(record)
        result["workbook_path"] = normalized
        return result

    @staticmethod
    @lru_cache(maxsize=128)
    def _find_cached(
        path_text: str,
        _modified_ns: int,
        _size_bytes: int,
        canonical_number: str,
    ) -> dict[str, Any]:
        workbook = load_workbook(
            Path(path_text),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        try:
            required_sheets = {OUTCOMES_SHEET, SUBSIDIES_SHEET}
            missing_sheets = required_sheets.difference(workbook.sheetnames)
            if missing_sheets:
                missing = ", ".join(sorted(missing_sheets))
                raise ValueError(f"Workbook is missing required sheets: {missing}")

            outcomes = workbook[OUTCOMES_SHEET]
            subsidies = workbook[SUBSIDIES_SHEET]
            outcome_header_row, outcome_headers = ProcessDataRepository._headers(
                outcomes,
                {PROCESS_COLUMN, "UF", "Sub-assunto", "Valor da causa"},
            )
            subsidy_header_row, subsidy_headers = ProcessDataRepository._headers(
                subsidies,
                {SUBSIDY_PROCESS_COLUMN, *BINARY_COLUMNS.values()},
            )
            outcome_row_number, outcome_row = ProcessDataRepository._find_unique_row(
                outcomes,
                outcome_header_row,
                outcome_headers[PROCESS_COLUMN],
                canonical_number,
            )
            subsidy_row_number, subsidy_row = ProcessDataRepository._find_unique_row(
                subsidies,
                subsidy_header_row,
                subsidy_headers[SUBSIDY_PROCESS_COLUMN],
                canonical_number,
            )

            stored_process_number = str(outcome_row[outcome_headers[PROCESS_COLUMN]]).strip()
            state = str(outcome_row[outcome_headers["UF"]]).strip().upper()
            if len(state) != 2:
                raise ValueError(f"Invalid UF for process {stored_process_number}")
            claim_amount = float(outcome_row[outcome_headers["Valor da causa"]])
            if claim_amount <= 0:
                raise ValueError(f"Invalid claim amount for process {stored_process_number}")

            return {
                "process_number": stored_process_number,
                "state": state,
                "sub_subject": ProcessDataRepository._sub_subject(
                    outcome_row[outcome_headers["Sub-assunto"]]
                ),
                "claim_amount": claim_amount,
                "evidence": {
                    field: ProcessDataRepository._binary(
                        subsidy_row[subsidy_headers[column]],
                        column,
                    )
                    for field, column in BINARY_COLUMNS.items()
                },
                "source_rows": {
                    OUTCOMES_SHEET: outcome_row_number,
                    SUBSIDIES_SHEET: subsidy_row_number,
                },
                "excluded_post_outcome_columns": list(EXCLUDED_POST_OUTCOME_COLUMNS),
            }
        finally:
            workbook.close()

    @staticmethod
    def _headers(sheet: Worksheet, required: set[str]) -> tuple[int, dict[str, int]]:
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=10, values_only=True),
            start=1,
        ):
            headers = {
                str(value).strip(): index
                for index, value in enumerate(row)
                if value is not None and str(value).strip()
            }
            if required.issubset(headers):
                return row_number, headers
        missing = ", ".join(sorted(required))
        raise ValueError(f"Could not find required workbook headers: {missing}")

    @staticmethod
    def _find_unique_row(
        sheet: Worksheet,
        header_row: int,
        process_column_index: int,
        canonical_number: str,
    ) -> tuple[int, tuple[Any, ...]]:
        match: tuple[int, tuple[Any, ...]] | None = None
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            value = row[process_column_index]
            if value is None:
                continue
            if canonical_process_number(str(value)) != canonical_number:
                continue
            if match is not None:
                raise ValueError(f"Duplicate process number in sheet '{sheet.title}'")
            match = row_number, row
        if match is None:
            raise ValueError(f"Process was not found in sheet '{sheet.title}'")
        return match

    @staticmethod
    def _sub_subject(value: object) -> str:
        normalized = unicodedata.normalize("NFKD", str(value).strip())
        normalized = "".join(
            character for character in normalized if not unicodedata.combining(character)
        )
        mapping = {"golpe": "fraud", "generico": "generic"}
        try:
            return mapping[normalized.casefold()]
        except KeyError as exc:
            raise ValueError(f"Unsupported sub-subject in workbook: {value}") from exc

    @staticmethod
    def _binary(value: object, column: str) -> bool:
        if value in (0, False):
            return False
        if value in (1, True):
            return True
        raise ValueError(f"Column '{column}' must contain 0 or 1")
