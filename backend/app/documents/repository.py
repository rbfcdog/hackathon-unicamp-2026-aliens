import csv
import hashlib
import re
import uuid
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, BinaryIO

import pymupdf
import pymupdf4llm
from openpyxl import load_workbook
from pymupdf4llm.ocr import OCRMode

SUPPORTED_DOCUMENT_SUFFIXES = frozenset({".pdf", ".csv", ".xlsx", ".xlsm"})
SUPPORTED_UPLOAD_SUFFIXES = frozenset({".pdf", ".csv"})
MAX_UPLOAD_BYTES = 20 * 1024 * 1024


@dataclass(frozen=True)
class SupportedDocument:
    path: str
    kind: str
    size_bytes: int

    def as_dict(self) -> dict[str, str | int]:
        return asdict(self)


class DocumentRepository:
    def __init__(self, root: str | Path) -> None:
        candidate = Path(root).expanduser()
        if not candidate.is_absolute():
            backend_root = Path(__file__).resolve().parents[2]
            candidate = backend_root / candidate
        self.root = candidate.resolve()
        if not self.root.is_dir():
            raise ValueError(f"Document root does not exist: {self.root}")

    def list_documents(self) -> list[SupportedDocument]:
        documents = []
        for path in self.root.rglob("*"):
            if not path.is_file() or path.suffix.lower() not in SUPPORTED_DOCUMENT_SUFFIXES:
                continue
            documents.append(
                SupportedDocument(
                    path=path.relative_to(self.root).as_posix(),
                    kind=self._kind(path),
                    size_bytes=path.stat().st_size,
                )
            )
        return sorted(documents, key=lambda document: document.path.casefold())

    def resolve(self, relative_path: str, *, expected_suffixes: set[str]) -> tuple[Path, str]:
        supplied = Path(relative_path)
        if supplied.is_absolute():
            raise ValueError("Document path must be relative to the configured document root")
        candidate = (self.root / supplied).resolve()
        try:
            normalized = candidate.relative_to(self.root).as_posix()
        except ValueError as exc:
            raise ValueError("Document path escapes the configured document root") from exc
        if candidate.suffix.lower() not in expected_suffixes:
            expected = ", ".join(sorted(expected_suffixes))
            raise ValueError(f"Unsupported document type; expected one of: {expected}")
        if not candidate.is_file():
            raise ValueError(f"Document does not exist: {normalized}")
        return candidate, normalized

    def store_upload(
        self,
        original_filename: str,
        source: BinaryIO,
        *,
        max_bytes: int = MAX_UPLOAD_BYTES,
    ) -> dict[str, str | int]:
        filename = original_filename.strip()
        if not filename or Path(filename).name != filename:
            raise ValueError("Upload filename must be a plain filename")
        if len(filename) > 200:
            raise ValueError("Upload filename must contain at most 200 characters")
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_UPLOAD_SUFFIXES:
            expected = ", ".join(sorted(SUPPORTED_UPLOAD_SUFFIXES))
            raise ValueError(f"Only these upload types are supported: {expected}")
        safe_filename = re.sub(r"[^A-Za-z0-9._-]+", "_", filename).strip("._")
        if not safe_filename or Path(safe_filename).suffix.lower() != suffix:
            safe_filename = f"document{suffix}"

        upload_id = str(uuid.uuid4())
        directory = self.root / "uploads" / upload_id
        directory.mkdir(parents=True, exist_ok=False)
        destination = directory / safe_filename
        temporary = directory / f".{safe_filename}.part"
        digest = hashlib.sha256()
        size_bytes = 0
        try:
            with temporary.open("wb") as target:
                while chunk := source.read(1024 * 1024):
                    size_bytes += len(chunk)
                    if size_bytes > max_bytes:
                        raise ValueError(f"Upload exceeds the {max_bytes}-byte limit")
                    digest.update(chunk)
                    target.write(chunk)
            if size_bytes == 0:
                raise ValueError("Uploaded document is empty")
            temporary.replace(destination)
            if suffix == ".pdf":
                self._validate_pdf_upload(destination)
            else:
                self._validate_csv_upload(destination)
        except Exception:
            temporary.unlink(missing_ok=True)
            destination.unlink(missing_ok=True)
            directory.rmdir()
            raise

        return {
            "upload_id": upload_id,
            "path": destination.relative_to(self.root).as_posix(),
            "kind": self._kind(destination),
            "size_bytes": size_bytes,
            "sha256": digest.hexdigest(),
        }

    def delete_upload(self, relative_path: str) -> None:
        path, normalized = self.resolve(
            relative_path,
            expected_suffixes=set(SUPPORTED_UPLOAD_SUFFIXES),
        )
        relative = Path(normalized)
        if len(relative.parts) != 3 or relative.parts[0] != "uploads":
            raise ValueError("Only managed uploads can be deleted")
        path.unlink()
        path.parent.rmdir()

    @staticmethod
    def _validate_pdf_upload(path: Path) -> None:
        with path.open("rb") as uploaded:
            if b"%PDF-" not in uploaded.read(1024):
                raise ValueError("Uploaded file does not contain a PDF header")
        try:
            with pymupdf.open(path) as document:
                if document.needs_pass:
                    raise ValueError("Encrypted PDF uploads are not supported")
                if document.page_count == 0:
                    raise ValueError("Uploaded PDF has no pages")
        except ValueError:
            raise
        except (pymupdf.FileDataError, RuntimeError) as exc:
            raise ValueError("Uploaded PDF is invalid") from exc

    @classmethod
    def _validate_csv_upload(cls, path: Path) -> None:
        encoding, dialect = cls._csv_format(path)
        with path.open("r", encoding=encoding, newline="") as uploaded:
            if not any(any(cell.strip() for cell in row) for row in csv.reader(uploaded, dialect)):
                raise ValueError("Uploaded CSV has no data")

    @staticmethod
    def _csv_format(path: Path) -> tuple[str, Any]:
        with path.open("rb") as source:
            sample_bytes = source.read(65_536)
        if b"\x00" in sample_bytes:
            raise ValueError("Uploaded CSV contains unsupported binary data")
        try:
            sample = sample_bytes.decode("utf-8-sig")
            encoding = "utf-8-sig"
        except UnicodeDecodeError:
            sample = sample_bytes.decode("cp1252")
            encoding = "cp1252"
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return encoding, dialect

    def read_pdf(
        self,
        relative_path: str,
        *,
        start_page: int = 1,
        max_pages: int | None = None,
        max_characters: int | None = None,
    ) -> dict[str, Any]:
        if start_page < 1:
            raise ValueError("start_page must be at least 1")
        if max_pages is not None and max_pages < 1:
            raise ValueError("max_pages must be at least 1")
        if max_characters is not None and max_characters < 1:
            raise ValueError("max_characters must be at least 1")

        path, normalized = self.resolve(relative_path, expected_suffixes={".pdf"})
        try:
            with pymupdf.open(path) as document:
                if document.needs_pass:
                    raise ValueError(f"Encrypted PDF is not supported: {normalized}")
                total_pages = document.page_count
        except ValueError:
            raise
        except (pymupdf.FileDataError, RuntimeError) as exc:
            raise ValueError(f"Invalid PDF: {normalized}") from exc
        if start_page > total_pages:
            raise ValueError(f"start_page exceeds the PDF page count ({total_pages})")

        last_page = (
            total_pages
            if max_pages is None
            else min(total_pages, start_page + max_pages - 1)
        )
        extracted_pages = pymupdf4llm.to_markdown(
            str(path),
            pages=list(range(start_page - 1, last_page)),
            page_chunks=True,
            use_ocr=OCRMode.SELECT_KEEP_OLD,
            ocr_dpi=300,
            ocr_language="por+eng",
            show_progress=False,
            table_output="markdown",
        )
        content_parts: list[str] = []
        characters = 0
        truncated = False
        pages_read = 0
        for page_number, page in enumerate(extracted_pages, start=start_page):
            pages_read += 1
            text = str(page["text"]).strip()
            if not text:
                continue
            chunk = f"--- página {page_number} ---\n{text}"
            part = f"\n\n{chunk}" if content_parts else chunk
            if max_characters is not None:
                remaining = max_characters - characters
                if len(part) > remaining:
                    content_parts.append(part[:remaining])
                    truncated = True
                    break
            content_parts.append(part)
            characters += len(part)

        content = "".join(content_parts).strip()
        return {
            "status": "ok" if content else "empty",
            "kind": "pdf",
            "document_path": normalized,
            "start_page": start_page,
            "end_page": start_page + pages_read - 1,
            "total_pages": total_pages,
            "content": content,
            "truncated": truncated or last_page < total_pages,
            "extraction_engine": "pymupdf4llm",
            "ocr_mode": "select_keep_old",
            "ocr_language": "por+eng",
            "ocr_dpi": 300,
        }

    def read_spreadsheet(
        self,
        relative_path: str,
        *,
        sheet_name: str | None = None,
        start_row: int = 1,
        max_rows: int | None = None,
        max_characters: int | None = None,
    ) -> dict[str, Any]:
        if start_row < 1:
            raise ValueError("start_row must be at least 1")
        if max_rows is not None and max_rows < 1:
            raise ValueError("max_rows must be at least 1")
        if max_characters is not None and max_characters < 1:
            raise ValueError("max_characters must be at least 1")

        path, normalized = self.resolve(
            relative_path,
            expected_suffixes={".xlsx", ".xlsm"},
        )
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            selected_sheet = sheet_name or workbook.sheetnames[0]
            if selected_sheet not in workbook.sheetnames:
                available = ", ".join(workbook.sheetnames)
                raise ValueError(f"Unknown sheet '{selected_sheet}'. Available sheets: {available}")
            sheet = workbook[selected_sheet]
            if sheet.max_row is None or sheet.max_column is None:
                sheet.calculate_dimension(force=True)
            total_rows = sheet.max_row or 0
            total_columns = sheet.max_column or 0
            if start_row > total_rows:
                raise ValueError(f"start_row exceeds the worksheet row count ({total_rows})")
            end_row = (
                total_rows
                if max_rows is None
                else min(total_rows, start_row + max_rows - 1)
            )
            lines: list[str] = []
            characters = 0
            truncated = False
            rows_read = 0
            for row_number, values in enumerate(
                sheet.iter_rows(
                    min_row=start_row,
                    max_row=end_row,
                    max_col=total_columns,
                    values_only=True,
                ),
                start=start_row,
            ):
                rendered = "\t".join(self._render_cell(value) for value in values).rstrip()
                line = f"[linha {row_number}] {rendered}\n"
                if max_characters is not None:
                    remaining = max_characters - characters
                    if len(line) > remaining:
                        lines.append(line[:remaining])
                        truncated = True
                        rows_read += 1
                        break
                lines.append(line)
                characters += len(line)
                rows_read += 1

            content = "".join(lines).strip()
            return {
                "status": "ok" if content else "empty",
                "kind": "spreadsheet",
                "document_path": normalized,
                "sheet_name": selected_sheet,
                "available_sheets": workbook.sheetnames,
                "start_row": start_row,
                "end_row": start_row + rows_read - 1,
                "total_rows": total_rows,
                "total_columns": total_columns,
                "content": content,
                "truncated": truncated or end_row < total_rows,
            }
        finally:
            workbook.close()

    def read_csv(
        self,
        relative_path: str,
        *,
        start_row: int = 1,
        max_rows: int | None = None,
        max_characters: int | None = None,
    ) -> dict[str, Any]:
        if start_row < 1:
            raise ValueError("start_row must be at least 1")
        if max_rows is not None and max_rows < 1:
            raise ValueError("max_rows must be at least 1")
        if max_characters is not None and max_characters < 1:
            raise ValueError("max_characters must be at least 1")

        path, normalized = self.resolve(relative_path, expected_suffixes={".csv"})
        encoding, dialect = self._csv_format(path)
        lines: list[str] = []
        characters = 0
        total_rows = 0
        total_columns = 0
        rows_read = 0
        content_truncated = False
        end_row = None if max_rows is None else start_row + max_rows - 1
        with path.open("r", encoding=encoding, newline="") as source:
            for row_number, values in enumerate(csv.reader(source, dialect), start=1):
                total_rows = row_number
                total_columns = max(total_columns, len(values))
                outside_window = row_number < start_row or (
                    end_row is not None and row_number > end_row
                )
                if outside_window or content_truncated:
                    continue
                rendered = "\t".join(self._render_cell(value) for value in values).rstrip()
                line = f"[linha {row_number}] {rendered}\n"
                if max_characters is not None:
                    remaining = max_characters - characters
                    if len(line) > remaining:
                        lines.append(line[:remaining])
                        content_truncated = True
                        rows_read += 1
                        continue
                lines.append(line)
                characters += len(line)
                rows_read += 1

        if start_row > total_rows:
            raise ValueError(f"start_row exceeds the CSV row count ({total_rows})")
        content = "".join(lines).strip()
        return {
            "status": "ok" if content else "empty",
            "kind": "csv",
            "document_path": normalized,
            "encoding": encoding,
            "delimiter": dialect.delimiter,
            "start_row": start_row,
            "end_row": start_row + rows_read - 1,
            "total_rows": total_rows,
            "total_columns": total_columns,
            "content": content,
            "truncated": content_truncated
            or (end_row is not None and end_row < total_rows),
        }

    @staticmethod
    def _render_cell(value: object) -> str:
        if value is None:
            return ""
        return str(value).replace("\r", " ").replace("\n", " ").replace("\t", " ")

    @staticmethod
    def _kind(path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return "pdf"
        if suffix == ".csv":
            return "csv"
        return "spreadsheet"
